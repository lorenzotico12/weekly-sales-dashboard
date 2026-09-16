
import streamlit as st
import pandas as pd
import os
import requests
from io import BytesIO
from pptx import Presentation
from pptx.util import Inches

st.set_page_config(page_title="Weekly Sales Dashboard", layout="wide")
st.title("📈 Weekly Sales Performance Dashboard")

st.markdown("Upload your 3 Excel files: **Current Week**, **Last Week**, and **Same Week Last Year**.\n")
uploaded_files = st.file_uploader("Upload Excel Files", type=["xlsx"], accept_multiple_files=True)

# Normalize and map column names
def standardize_columns(df):
    df.columns = [col.strip().lower() for col in df.columns]
    col_map = {
        "sold itmens after return": "sold items after return",
        "sold items after return": "sold items after return",
        "return rate": "return rate",
        "cg2": "cg2",
        "cg3": "cg3",
        "supplier article name": "supplier article name",
        "image link": "image link"
    }
    return df.rename(columns={col: col_map.get(col, col) for col in df.columns})

def generate_summary_and_images(current_df, last_df, year_df):
    kpi_lines = []

    current_sales = current_df["sold items after return"].sum()
    last_sales = last_df["sold items after return"].sum()
    year_sales = year_df["sold items after return"].sum()

    wow_change = ((current_sales - last_sales) / last_sales) * 100 if last_sales else 0
    yoy_change = ((current_sales - year_sales) / year_sales) * 100 if year_sales else 0

    kpi_lines.append(f"• Sales: {int(current_sales)} units ({wow_change:+.1f}% WoW, {yoy_change:+.1f}% YoY)")

    current_return = current_df["return rate"].mean()
    last_return = last_df["return rate"].mean()
    year_return = year_df["return rate"].mean()

    kpi_lines.append(f"• Return Rate: {current_return:.1f}% ({(current_return - last_return):+.1f} pts WoW, {(current_return - year_return):+.1f} pts YoY)")

    gender_data = current_df.groupby("cg2")["sold items after return"].sum()
    total = gender_data.sum()
    for gender, count in gender_data.items():
        percent = (count / total) * 100
        kpi_lines.append(f"• {gender}: {int(count)} units ({percent:.1f}%)")

    return kpi_lines

def add_top5_slide(prs, df):
    slide_layout = prs.slide_layouts[5]
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "🏆 Top 5 Best Sellers – Current Week"

    top5 = df.sort_values("sold items after return", ascending=False).head(5)

    left = Inches(0.5)
    top = Inches(1.0)
    width = Inches(1.5)
    height = Inches(1.5)

    for i, (_, row) in enumerate(top5.iterrows()):
        x = left + (i % 5) * Inches(2)
        y = top

        try:
            response = requests.get(row["image link"])
            image_stream = BytesIO(response.content)
            slide.shapes.add_picture(image_stream, x, y, width, height)
        except Exception:
            slide.shapes.add_textbox(x, y, width, height).text = "Image not found"

        details = f"{row['supplier article name']}\n{row['cg2']} / {row['cg3']}\n{int(row['sold items after return'])} units"
        text_box = slide.shapes.add_textbox(x, y + Inches(1.6), width, Inches(1))
        frame = text_box.text_frame
        for line in details.split("\n"):
            frame.add_paragraph().text = line

def export_to_powerpoint(kpis, df):
    prs = Presentation()
    slide_layout = prs.slide_layouts[5]
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "📊 Weekly KPI Summary"

    text_box = slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(8), Inches(5))
    frame = text_box.text_frame
    for line in kpis:
        frame.add_paragraph().text = line

    add_top5_slide(prs, df)

    os.makedirs("output", exist_ok=True)
    ppt_path = "output/Weekly_Sales_Report.pptx"
    prs.save(ppt_path)
    return ppt_path

if uploaded_files and len(uploaded_files) == 3:
    labels = ["Current Week", "Last Week", "Same Week Last Year"]
    dataframes = [standardize_columns(pd.read_excel(file)) for file in uploaded_files]

    for label, df in zip(labels, dataframes):
        st.subheader(f"{label} Summary")
        st.dataframe(df.head())

    if st.button("Generate KPI Report and PowerPoint"):
        kpis = generate_summary_and_images(*dataframes)
        ppt_path = export_to_powerpoint(kpis, dataframes[0])
        st.success("PowerPoint generated successfully!")
        with open(ppt_path, "rb") as f:
            st.download_button("📥 Download PowerPoint", f, file_name="Weekly_Sales_Report.pptx")

else:
    st.info("Please upload 3 Excel files.")
